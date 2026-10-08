# INSIGHT. 팀 QA 시트 생성 — 사용: python scripts/qa/make_qa_sheet.py . ../QA/INSIGHT_QA_시트.xlsx  (openpyxl 필요, .env.local 의 Supabase 키로 글 목록을 읽는다)
import json, re, sqlite3, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

PROJECT, OUT = sys.argv[1], sys.argv[2]

INK, SUB, LINE = "191F28", "6B7684", "E5E8EB"
BRAND, BRAND_SOFT = "3182F6", "E8F3FF"
HEAD = PatternFill("solid", fgColor="191F28")
SECTION = PatternFill("solid", fgColor="F2F4F6")
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
F = "맑은 고딕"

def font(**k):
    return Font(name=F, size=k.pop("size", 10), **k)

wb = Workbook()

# ───────────── 공용 ─────────────
RESULTS = ["미진행", "Pass", "Fail", "Block", "N/A"]
SEVERITY = ["S1 치명", "S2 높음", "S3 보통", "S4 낮음"]
ISSUE_TYPES = ["기능 오류", "UI·레이아웃", "문구·오탈자", "AI 정리 내용", "경험 분류", "성능·속도", "개선 제안"]
STATUS = ["신규", "확인 중", "수정 중", "수정 완료", "재확인 완료", "보류", "이슈 아님"]
ENVS = ["PC · Chrome", "PC · Edge", "PC · Safari", "iPhone · Safari", "Android · Chrome", "기타"]
SCREENS = ["공통(상단·탭바)", "홈", "경험 페이지", "글 상세", "피드", "검색", "기술 페이지", "마이", "제외된 글"]
EXP_SHORT = {
    "WAIT_LESS": "기다림 없는 경험", "SEARCH_LESS": "찾지 않아도 보이는 경험", "COMPARE_LESS": "고민을 덜어 주는 경험",
    "INPUT_LESS": "다시 입력하지 않는 경험", "REPEAT_LESS": "나를 기억해 주는 경험", "DO_LESS": "알아서 처리되는 경험",
    "CHECK_LESS": "확인하지 않아도 안심되는 경험",
}

def dv(ws, options, rng):
    d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showDropDown=False)
    d.error, d.errorTitle = "목록에서 골라 주세요.", "입력 값 확인"
    ws.add_data_validation(d)
    d.add(rng)

def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = font(bold=True, color="FFFFFF"), HEAD, CENTER, BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def color_rule(ws, rng, first_cell, text, fill, color):
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'{first_cell}="{text}"'], fill=PatternFill("solid", fgColor=fill), font=Font(name=F, color=color, bold=True))
    )

def result_colors(ws, rng, first):
    color_rule(ws, rng, first, "Pass", "E3F6EA", "1B7F3B")
    color_rule(ws, rng, first, "Fail", "FDE7E7", "C62828")
    color_rule(ws, rng, first, "Block", "FFF2D9", "A15C00")
    color_rule(ws, rng, first, "N/A", "F2F4F6", "6B7684")

# ═════════════ 1. 안내 ═════════════
ws = wb.active
ws.title = "안내"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 62
ws.column_dimensions["D"].width = 40

r = 2
ws.cell(row=r, column=2, value="INSIGHT. 팀 QA").font = font(size=20, bold=True, color=INK)
r += 1
ws.cell(row=r, column=2, value="테크 블로그 관점 아카이브 — 배포 전 팀 QA 진행 기준과 기록 시트").font = font(size=11, color=SUB)
r += 2

def section(title):
    global r
    for c in range(2, 5):
        ws.cell(row=r, column=c).fill = SECTION
    ws.cell(row=r, column=2, value=title).font = font(size=12, bold=True, color=INK)
    ws.row_dimensions[r].height = 22
    r += 1

def kv(k, v, note="", input_cell=False):
    global r
    ws.cell(row=r, column=2, value=k).font = font(bold=True, color=INK)
    c = ws.cell(row=r, column=3, value=v)
    c.font = font(color=BRAND if input_cell else INK, bold=input_cell)
    c.alignment = WRAP
    if input_cell:
        c.fill = PatternFill("solid", fgColor=BRAND_SOFT)
        c.border = BOX
    if note:
        n = ws.cell(row=r, column=4, value=note)
        n.font, n.alignment = font(size=9, color=SUB), WRAP
    for cc in (2,):
        ws.cell(row=r, column=cc).alignment = WRAP
    r += 1

section("1. 기본 정보  (파란 칸은 직접 채워 주세요)")
kv("배포 주소", "https://", "배포 후 주소를 넣으면 '콘텐츠 검수' 시트의 [열기] 링크가 이 주소로 열려요. 끝에 / 는 빼 주세요.", True)
BASE_URL_CELL = f"안내!$C${r-1}"
kv("QA 기간", "2026.10.  ~  2026.10.", "", True)
kv("참여자", "", "이름을 쉼표로 적어 주세요.", True)
kv("QA 담당(취합)", "", "이슈 우선순위를 정하고 수정 담당을 배정하는 사람", True)
kv("기준 데이터", "분류 v4 · Supabase 기준 (2026-10-08)", "QA 중 배포가 바뀌면 여기와 이슈의 '비고'에 적어 주세요.")
r += 1

section("2. 무엇을 확인하나요")
kv("① 기능·화면", "'테스트 케이스' 시트 — 화면별로 정해 둔 항목을 PC와 모바일에서 한 번씩 확인해요.")
kv("② 콘텐츠 품질", "'콘텐츠 검수' 시트 — AI가 정리한 내용과 경험 분류가 원문과 맞는지 글마다 확인해요. 이 서비스에서 가장 중요한 QA예요.")
kv("③ 제외 판정", "'제외 글 검수' 시트 — 기획자에게 필요 없다고 뺀 글이 정말 빼도 되는지 확인해요.")
kv("④ 이슈 기록", "①~③에서 문제를 찾으면 '이슈 리포트' 시트에 한 줄로 남기고, 해당 칸에 이슈 ID를 적어요.")
kv("⑤ 현황", "'현황' 시트에서 진행률과 심각도별 이슈 수가 자동으로 집계돼요.")
r += 1

section("3. 진행 순서")
steps = [
    "자기 이름으로 담당 범위를 나눠요 — 테스트 케이스는 화면 단위, 콘텐츠 검수는 글 번호 구간 단위로 나누면 겹치지 않아요.",
    "시크릿 창(또는 새 브라우저)으로 시작해요. 읽음·저장·최근 검색은 브라우저에만 저장돼서, 이전 기록이 있으면 결과가 달라져요.",
    "테스트 케이스를 PC → 모바일 순으로 확인하고 결과 칸에 Pass / Fail / Block / N/A 를 골라요.",
    "Fail 이나 Block 이면 '이슈 리포트'에 새 줄을 쓰고, 테스트 케이스의 '이슈 ID' 칸에 같은 번호를 적어요.",
    "콘텐츠 검수는 [열기]로 글 상세를 보고, 필요하면 [원문]과 비교해서 판단해요. 애매하면 '△'와 의견을 남겨 주세요.",
    "수정이 끝난 이슈는 등록한 사람이 다시 확인하고 상태를 '재확인 완료'로 바꿔요.",
]
for i, s in enumerate(steps, 1):
    kv(f"STEP {i}", s)
r += 1

section("4. 결과 값")
for k, v in [
    ("Pass", "기대 결과대로 동작해요."),
    ("Fail", "기대 결과와 달라요. → 이슈 리포트 작성"),
    ("Block", "다른 문제 때문에 확인할 수 없어요(예: 화면이 안 열림). → 막은 이슈 ID 기록"),
    ("N/A", "이 환경에는 해당하지 않아요(예: PC 전용 화살표를 모바일에서)."),
    ("미진행", "아직 확인하지 않았어요. 기본 값이에요."),
]:
    kv(k, v)
r += 1

section("5. 심각도 기준")
for k, v, ex in [
    ("S1 치명", "서비스를 쓸 수 없어요. 우회 방법이 없어요.", "화면이 열리지 않음(500·빈 화면), 모든 글 상세가 깨짐, 앱이 멈춤"),
    ("S2 높음", "핵심 기능이 잘못 동작해요. 사용자가 잘못된 정보를 보게 돼요.", "필터 결과가 틀림, 다른 글로 이동, 원문과 반대되는 AI 요약, 분류가 명백히 틀림"),
    ("S3 보통", "기능은 되지만 불편하거나 일부 기기에서만 깨져요. 우회할 수 있어요.", "모바일에서 카드가 잘림, 넘김이 매끄럽지 않음, 하이라이트 위치가 어긋남"),
    ("S4 낮음", "작은 문제나 개선 제안이에요.", "오탈자, 여백·정렬, 문구 다듬기, 이런 기능이 있으면 좋겠어요"),
]:
    kv(k, v, ex)
kv("배포 기준", "S1·S2 가 0건이면 배포해요. S3 은 배포 후 바로 고치고, S4 는 다음 개선에 모아요.")
r += 1

section("6. 이슈 쓰는 법")
for k, v in [
    ("한 줄 = 문제 하나", "여러 문제를 한 줄에 묶지 마세요. 따로 고치고 따로 확인해요."),
    ("제목", "어디서 · 무엇이 · 어떻게 — 예) 글 상세 · 기획자 관점 시트 · 모바일에서 닫기 버튼이 안 눌려요"),
    ("재현 단계", "1. 2. 3. 번호를 붙여 누구나 따라 할 수 있게 써요. 시작 주소를 꼭 적어 주세요."),
    ("기대 / 실제", "어떻게 될 줄 알았는지, 실제로 어떻게 됐는지 나눠서 써요."),
    ("스크린샷", "구글 드라이브 공유 폴더에 올리고 링크를 붙여요. 움직임 문제는 화면 녹화가 좋아요."),
    ("AI 내용 문제", "유형을 'AI 정리 내용' 또는 '경험 분류'로 고르고, 원문의 어느 문단과 다른지 적어 주세요."),
]:
    kv(k, v)

# ═════════════ 2. 테스트 케이스 ═════════════
tc = wb.create_sheet("테스트 케이스")
cols = ["ID", "화면", "기능", "확인 항목", "확인 방법", "기대 결과", "우선순위", "담당", "PC 결과", "모바일 결과", "이슈 ID", "비고"]
header(tc, 1, cols, [8, 13, 16, 34, 44, 44, 9, 10, 11, 11, 10, 24])
tc.freeze_panes = "E2"

CASES = [
    # 화면, 기능, 확인 항목, 확인 방법, 기대 결과, 우선순위
    ("공통(상단·탭바)", "상단 바", "로고·메뉴가 보이나요", "아무 화면에서 맨 위를 봐요", "INSIGHT. 로고(점만 파란색)와 홈·피드·검색·마이 메뉴가 보이고, 로고를 누르면 홈으로 가요", "P1"),
    ("공통(상단·탭바)", "하단 탭바", "모바일 탭 이동", "모바일에서 아래 탭 4개를 차례로 눌러요", "각 화면으로 이동하고 지금 화면의 탭이 파란색이에요. 아이콘이 이모지가 아닌 선 아이콘이에요", "P1"),
    ("공통(상단·탭바)", "내가 읽은 글 토글", "홈에서만 보이나요", "홈과 다른 화면의 오른쪽 위를 비교해요", "'내가 읽은 글' 스위치는 홈에서만 보여요", "P2"),
    ("공통(상단·탭바)", "새로고침·직접 접속", "주소로 바로 들어가기", "글 상세·경험·기술 페이지 주소를 새 탭에 붙여 넣어요", "오류 없이 같은 화면이 열려요", "P1"),
    ("공통(상단·탭바)", "없는 주소", "404 처리", "/articles/99999 처럼 없는 주소로 들어가요", "오류 화면 대신 '찾을 수 없음' 안내가 나와요", "P3"),
    ("홈", "새 글 카드", "카드 수와 자동 넘김", "홈에 들어와 가만히 기다려요", "최근 글 최대 4장이 약 4.5초마다 넘어가고, 마지막 다음은 첫 장으로 돌아와요", "P1"),
    ("홈", "새 글 카드", "손을 대면 멈춤", "카드에 마우스를 올리거나 손가락으로 누르고 있어요", "자동 넘김이 멈추고, 떼면 다시 넘어가요", "P3"),
    ("홈", "새 글 카드", "썸네일이 없는 글", "썸네일이 없는 카드를 찾아봐요", "회사 로고와 회사 색이 흐릿하게 깔린 배경이 보여요(빈 회색 칸이 아니에요)", "P2"),
    ("홈", "새 글 카드", "읽으러 가기", "'읽으러 가기'를 눌러요", "그 카드의 글 상세로 가요", "P1"),
    ("홈", "내가 읽은 글 토글", "OFF 이면 읽은 글이 빠지나요", "글 2개를 읽고 홈으로 돌아와 스위치를 꺼요", "새 글 카드·문제 카드·기술 목록에서 읽은 글이 빠지고, 안 읽은 글로 다시 채워져요", "P1"),
    ("홈", "내가 읽은 글 토글", "ON 이면 읽음 표시", "스위치를 다시 켜요", "읽은 글이 다시 보이고 '✓ 읽음' 표시가 붙어요. 새로고침해도 마지막 선택이 유지돼요", "P2"),
    ("홈", "경험 카드", "카드 구성", "'우리 사용자에게 어떤 경험을 주고 싶나요?' 카드를 봐요", "질문 → 경험 이름 → 파란 질문 → 대표 사례 → '다른 서비스 사례 N건 보기' 순서예요. 이모지가 없어요", "P1"),
    ("홈", "경험 카드", "넘기기와 위치 점", "카드를 옆으로 넘기고 아래 점을 눌러 봐요", "손가락·점 모두로 넘어가고, 지금 위치의 점이 길게 표시돼요", "P2"),
    ("홈", "경험 카드", "대표 사례·버튼 이동", "대표 사례와 '다른 서비스 사례 N건 보기'를 눌러요", "대표 사례는 글 상세로, 버튼은 그 경험 페이지로 가요. 버튼의 건수와 경험 페이지의 '전체' 건수가 같아요", "P1"),
    ("홈", "문제 카드", "PC 화살표", "PC에서 '이런 문제를 어떻게 풀었을까요?' 오른쪽 › 를 눌러요", "끝까지 넘어가고 마지막 카드가 오른쪽 끝에 맞게 멈춰요. 처음에는 ‹ 가, 끝에서는 › 가 사라져요", "P1"),
    ("홈", "문제 카드", "PC 마우스 끌기", "카드를 마우스로 잡고 옆으로 끌어요", "카드가 따라 넘어가고, 끌고 놓았을 때 글이 열리지 않아요", "P2"),
    ("홈", "문제 카드", "모바일 넘기기", "모바일에서 손가락으로 넘겨요", "마지막 카드까지 닿고, 첫 카드 왼쪽에 여백이 있어요", "P1"),
    ("홈", "문제 카드", "읽기 쉬운 색", "노랑·연두처럼 밝은 회사 색 카드를 봐요", "밝은 배경에서는 글자가 어두운 색이라 잘 읽혀요", "P3"),
    ("홈", "요즘 자주 나오는 기술", "탭 전환", "기술 키워드를 하나씩 눌러요", "아래 글 목록이 그 기술의 글로 바뀌고, '# … 글 N건 모두 보기'로 기술 페이지에 가요", "P1"),
    ("홈", "전체 글 보기", "피드로 이동", "맨 아래 '전체 글 보기'를 눌러요", "피드로 가요", "P3"),
    ("경험 페이지", "머리 영역", "홈 카드와 같은 문구", "홈 경험 카드에서 들어와요", "질문·경험 이름·파란 질문이 홈 카드와 같아요", "P2"),
    ("경험 페이지", "누구의 경험 필터", "전체 / 서비스 사용자 / 사내 운영·개발", "칩을 차례로 눌러요", "목록이 바뀌고 칩 숫자와 실제 카드 수가 맞아요", "P1"),
    ("경험 페이지", "읽음 필터", "전체 / 안 읽은 글 / 읽은 글", "글을 하나 읽고 돌아와 필터를 바꿔요", "읽은 글만·안 읽은 글만 보여요. 다 읽었거나 없으면 안내 문구가 나와요", "P1"),
    ("경험 페이지", "방식별 묶음", "묶음 제목과 건수", "아래로 내려 묶음을 봐요", "방식 이름 아래 건수가 카드 수와 같아요. '함께 볼 글'은 맨 아래에 따로 있어요", "P2"),
    ("경험 페이지", "다른 경험", "아래 경험 칩", "맨 아래 다른 경험 칩을 눌러요", "그 경험 페이지로 가요", "P3"),
    ("글 상세", "읽는 순서", "위에서 아래로 흐름", "아무 글이나 열어 처음부터 읽어 내려가요", "회사 → 원제목(작은 회색) → 핵심 카피 → AI 정리 안내 → 핵심 내용 → 결론부터 말하면(BEFORE → AFTER → 경험 칩) 순서예요", "P1"),
    ("글 상세", "회사 줄", "회사 로고·발행일·링크", "맨 위 회사 줄을 봐요·눌러요", "로고와 '○○ 기술 블로그', 발행일이 보이고, 누르면 그 회사 블로그가 새 탭으로 열려요", "P2"),
    ("글 상세", "읽음 처리", "열면 자동으로 읽음", "안 읽은 글을 열고 홈·피드로 돌아가요", "그 글에 '✓ 읽음' 표시가 붙어요", "P1"),
    ("글 상세", "읽음 토글", "직접 바꾸기", "오른쪽 위 '✓ 읽음' 버튼을 눌러요", "'안 읽음'으로 바뀌고 목록에서도 읽음 표시가 사라져요", "P2"),
    ("글 상세", "저장", "저장·해제", "'☆ 저장'을 눌렀다가 마이에서 확인하고 다시 해제해요", "마이에 들어갔다가, 해제하면 빠져요", "P1"),
    ("글 상세", "결론부터 말하면", "BEFORE / AFTER", "결론 아래 상자를 봐요", "회색 BEFORE → 화살표 → 초록 AFTER 순서이고, 아래에 경험 칩이 하나 있어요(해당 없음 글은 칩 없음)", "P1"),
    ("글 상세", "기획자 관점 보기", "바텀시트 열고 닫기", "'✦ 기획자 관점 보기'를 누르고 × 와 바깥을 눌러 닫아요", "적용 관점·적용 조건·토론 질문이 보이고, 핵심 변화(BEFORE/AFTER)는 중복으로 나오지 않아요. 닫으면 원래 위치예요", "P1"),
    ("글 상세", "한눈에 보기", "네 칸 요약", "'한눈에 보기'를 봐요", "무엇을·왜·어떻게·그래서 네 칸과 기획 포인트가 보여요", "P1"),
    ("글 상세", "더 들어가 볼까요?", "질문 칸 펼치기", "질문 칸을 하나씩 펼쳐요", "문제 → 해결 → 결말이 보이고, '근거 문장 N곳 보기'로 짧은 발췌(120자 안)와 '원문에서 이어 읽기 ↗'가 보여요", "P1"),
    ("글 상세", "더 들어가 볼까요?", "글 끝까지 다루나요", "긴 글에서 마지막 질문 칸이 원문의 어느 부분을 다루는지 원문과 비교해요", "질문 칸이 글 앞쪽에만 몰려 있지 않고 뒷부분 내용도 다뤄요", "P2"),
    ("글 상세", "원문에서 짚어 볼 문장", "핵심 문장 카드", "아래쪽 '원문에서 짚어 볼 문장'을 봐요", "종류(문제 정의·기술 이해·경험 변화) 표시, 짧은 인용 문장, 설명이 보여요. 인용 문장은 원문에 실제로 있는 문장이에요", "P2"),
    ("글 상세", "원문 비노출", "원문 전체가 보이지 않나요", "글 상세를 끝까지 내려 봐요", "원문 전체(문단 그대로)가 어디에도 보이지 않고, 짧은 발췌와 '원문에서 보기 ↗'만 있어요", "P1"),
    ("글 상세", "원문에서 보기", "원문 링크", "맨 아래 '원문에서 보기 ↗'를 눌러요", "원래 블로그 글이 새 탭으로 열려요(회사 블로그 이름이 버튼에 보여요)", "P1"),
    ("글 상세", "분석 전 글", "안내 문구", "(있다면) 분석 대기 글을 열어요", "'아직 AI 분석 전인 글이에요' 안내와 원문만 보여요", "P3"),
    ("피드", "목록", "최신순", "피드를 내려요", "기획 관점이 있는 글만 최신순으로 보여요", "P1"),
    ("피드", "회사 필터", "회사로 거르기", "회사 칩을 하나 눌러요", "그 회사 글만 보이고 다시 누르면 풀려요", "P1"),
    ("피드", "태그 필터", "태그로 거르기", "태그 칩을 눌러요", "그 태그 글만 보여요. 회사와 함께 고르면 둘 다 맞는 글만 보여요", "P2"),
    ("피드", "읽음 필터", "전체 / 안 읽은 글 / 읽은 글", "필터를 바꿔요", "숫자와 카드 수가 맞아요", "P2"),
    ("검색", "검색", "키워드 검색", "'검색' 'AI' 회사 이름 등으로 찾아봐요", "제목·회사·요약·가이드에 그 말이 있는 글이 나와요", "P1"),
    ("검색", "결과 없음", "안내 문구", "'ㅁㄴㅇㄹ'처럼 없는 말로 찾아요", "결과 없음 안내가 나와요", "P3"),
    ("검색", "최근 검색", "기록과 지우기", "몇 번 검색한 뒤 다시 검색 화면으로 와요", "최근 검색어가 보이고 눌러서 다시 찾거나 지울 수 있어요", "P2"),
    ("검색", "추천 관점·태그", "눌러서 찾기", "'이런 관점은 어때요?'와 태그를 눌러요", "그 말로 검색돼요", "P3"),
    ("기술 페이지", "기술별 글", "목록과 다른 기술", "홈 기술 탭에서 '모두 보기'로 들어와요", "그 기술 글이 보이고, 위 칩으로 다른 기술로 옮겨 갈 수 있어요", "P2"),
    ("마이", "저장한 글·읽은 글", "목록", "몇 개를 저장·읽은 뒤 마이를 열어요", "저장한 글과 읽은 글이 나뉘어 보여요. 비어 있으면 안내 문구가 나와요", "P1"),
    ("마이", "브라우저 저장 안내", "다른 기기", "다른 브라우저로 마이를 열어요", "기록이 없어요(이 브라우저에만 저장된다는 안내와 같아요)", "P3"),
    ("제외된 글", "목록", "제외 이유", "/excluded 로 들어가요", "제외된 글과 제외 이유 한 줄이 보여요", "P2"),
    ("공통(상단·탭바)", "반응형", "화면 폭 바꾸기", "PC 브라우저 폭을 줄였다 늘려요", "가로 스크롤이 생기지 않고, 좁아지면 아래 탭바로 바뀌어요", "P2"),
    ("공통(상단·탭바)", "속도", "첫 화면", "새로고침 후 홈이 다 보일 때까지 봐요", "3초 안에 내용이 보여요. 오래 걸리면 걸린 시간과 네트워크를 적어 주세요", "P2"),
    ("공통(상단·탭바)", "폰트·이모지", "글꼴과 이모지", "여러 화면을 둘러봐요", "모든 글자가 Pretendard 이고, 화면에 이모지가 없어요", "P3"),
    ("공통(상단·탭바)", "키보드", "Tab 으로 이동", "PC에서 Tab 키로 버튼을 옮겨 다녀요", "지금 위치가 보이고 Enter 로 눌려요", "P3"),
]
prefix = {"공통(상단·탭바)": "CM", "홈": "HM", "경험 페이지": "EX", "글 상세": "AR", "피드": "FD", "검색": "SR", "기술 페이지": "TC", "마이": "MY", "제외된 글": "XC"}
counter = {}
for i, (scr, feat, item, how, exp, pri) in enumerate(CASES, start=2):
    counter[scr] = counter.get(scr, 0) + 1
    vals = [f"{prefix[scr]}-{counter[scr]:02d}", scr, feat, item, how, exp, pri, "", "미진행", "미진행", "", ""]
    for c, v in enumerate(vals, start=1):
        cell = tc.cell(row=i, column=c, value=v)
        cell.font, cell.border = font(color=INK), BOX
        cell.alignment = CENTER if c in (1, 7, 9, 10, 11) else WRAP
last = len(CASES) + 1
dv(tc, RESULTS, f"I2:J{last}")
dv(tc, ["P1", "P2", "P3"], f"G2:G{last}")
result_colors(tc, f"I2:J{last}", "I2")
color_rule(tc, f"G2:G{last}", "G2", "P1", "FDE7E7", "C62828")
tc.auto_filter.ref = f"A1:L{last}"

# ═════════════ 3. 콘텐츠 검수 ═════════════
import urllib.request, urllib.parse
ENV = dict(l.split("=", 1) for l in open(f"{PROJECT}/.env.local", encoding="utf-8").read().splitlines() if "=" in l and not l.startswith("#"))
def sb_rows(status, cols):
    q = urllib.parse.urlencode({"select": cols, "status": f"eq.{status}", "order": "published_at.desc"})
    req = urllib.request.Request(f"{ENV['SUPABASE_URL']}/rest/v1/articles?{q}", headers={"apikey": ENV["SUPABASE_SERVICE_ROLE_KEY"], "Authorization": f"Bearer {ENV['SUPABASE_SERVICE_ROLE_KEY']}"})
    return json.load(urllib.request.urlopen(req))
src = open(f"{PROJECT}/src/core/0-collect/companies.ts", encoding="utf-8").read()
COMPANY = dict(re.findall(r'id: "([^"]+)",\s*name: "([^"]+)"', src))

cs = wb.create_sheet("콘텐츠 검수")
cols = ["글 ID", "회사", "원제목", "핵심 카피 (AI)", "AI 경험 분류", "적합도", "열기", "원문", "담당",
        "핵심 카피·결론 정확", "BEFORE→AFTER 정확", "더 들어가 볼까요 품질", "경험 분류 동의", "맞다고 보는 경험", "의견", "이슈 ID"]
header(cs, 1, cols, [7, 11, 30, 36, 20, 8, 7, 7, 9, 12, 12, 12, 13, 22, 30, 9])
cs.freeze_panes = "E2"
rows = [(r["id"], r["company_id"], r["title"], r["url"], json.dumps(r["learning"]) if r["learning"] else None) for r in sb_rows("included", "id,company_id,title,url,learning")]
for i, (aid, cid, title, url, lj) in enumerate(rows, start=2):
    l = json.loads(lj) if lj else {}
    c = l.get("classification") or {}
    exp = EXP_SHORT.get(c.get("experience"), "해당 없음" + (f" ({c.get('articleType')})" if c.get("articleType") else ""))
    vals = [aid, COMPANY.get(cid, cid), title, l.get("hook", ""), exp, c.get("fit", ""),
            f'=HYPERLINK({BASE_URL_CELL}&"/articles/{aid}","열기")', "원문",
            "", "", "", "", "", "", "", ""]
    for col, v in enumerate(vals, start=1):
        cell = cs.cell(row=i, column=col, value=v)
        cell.font, cell.border = font(color=INK), BOX
        cell.alignment = CENTER if col in (1, 6, 7, 8, 10, 11, 12, 13) else WRAP
        if col in (7, 8):
            cell.font = font(color=BRAND, underline="single")
        if col == 8:
            cell.hyperlink = url  # 수식 문자열은 255자 제한이 있어 긴 원문 주소는 셀 링크로 단다
last = len(rows) + 1
dv(cs, ["O 정확", "△ 일부 어긋남", "X 틀림"], f"J2:L{last}")
dv(cs, ["동의", "다른 경험이 맞음", "해당 없음이 맞음", "애매함"], f"M2:M{last}")
dv(cs, list(EXP_SHORT.values()) + ["해당 없음"], f"N2:N{last}")
for col in ("J", "K", "L"):
    rng = f"{col}2:{col}{last}"
    color_rule(cs, rng, f"{col}2", "O 정확", "E3F6EA", "1B7F3B")
    color_rule(cs, rng, f"{col}2", "△ 일부 어긋남", "FFF2D9", "A15C00")
    color_rule(cs, rng, f"{col}2", "X 틀림", "FDE7E7", "C62828")
color_rule(cs, f"M2:M{last}", "M2", "동의", "E3F6EA", "1B7F3B")
color_rule(cs, f"M2:M{last}", "M2", "다른 경험이 맞음", "FDE7E7", "C62828")
color_rule(cs, f"M2:M{last}", "M2", "해당 없음이 맞음", "FDE7E7", "C62828")
cs.auto_filter.ref = f"A1:P{last}"
n_content = len(rows)

# 판단 기준 메모(맨 아래)
guide_row = last + 2
for k, v in [
    ("판단 기준", ""),
    ("핵심 카피·결론 정확", "원문이 말하지 않은 내용·숫자가 없나요? 결론이 원문의 결론과 같은 방향인가요?"),
    ("BEFORE→AFTER 정확", "바뀌기 전과 후가 원문에 실제로 있는 장면인가요? 누가 무엇을 덜 하게 됐는지 보이나요?"),
    ("더 들어가 볼까요 품질", "질문 칸이 글 끝까지 고르게 다루나요? 개발 용어 없이 읽히나요? 칸끼리 같은 말을 반복하지 않나요?"),
    ("경험 분류 동의", "이 글이 '사용자(또는 사내 운영·개발자)가 무엇을 덜 하게 됐는지'를 보여 주나요? 그게 이 경험이 맞나요? 아니면 오른쪽에 맞다고 보는 경험을 골라 주세요."),
]:
    a = cs.cell(row=guide_row, column=3, value=k)
    a.font = font(bold=True, color=INK)
    b = cs.cell(row=guide_row, column=4, value=v)
    b.font, b.alignment = font(color=SUB), WRAP
    cs.merge_cells(start_row=guide_row, start_column=4, end_row=guide_row, end_column=8)
    cs.row_dimensions[guide_row].height = 30 if v else 18
    guide_row += 1

# ═════════════ 4. 제외 글 검수 ═════════════
xs = wb.create_sheet("제외 글 검수")
cols = ["글 ID", "회사", "원제목", "제외 이유 (AI)", "원문", "담당", "제외가 맞나요?", "의견"]
header(xs, 1, cols, [7, 11, 40, 60, 7, 9, 16, 36])
xs.freeze_panes = "D2"
rows = [(r["id"], r["company_id"], r["title"], r["url"], r["exclusion_reason"]) for r in sb_rows("excluded", "id,company_id,title,url,exclusion_reason")]
for i, (aid, cid, title, url, reason) in enumerate(rows, start=2):
    vals = [aid, COMPANY.get(cid, cid), title, reason or "", "원문", "", "", ""]
    for col, v in enumerate(vals, start=1):
        cell = xs.cell(row=i, column=col, value=v)
        cell.font, cell.border = font(color=INK), BOX
        cell.alignment = CENTER if col in (1, 5, 7) else WRAP
        if col == 5:
            cell.font = font(color=BRAND, underline="single")
            cell.hyperlink = url
last = len(rows) + 1
dv(xs, ["맞아요", "포함해야 해요", "애매해요"], f"G2:G{last}")
color_rule(xs, f"G2:G{last}", "G2", "맞아요", "E3F6EA", "1B7F3B")
color_rule(xs, f"G2:G{last}", "G2", "포함해야 해요", "FDE7E7", "C62828")
color_rule(xs, f"G2:G{last}", "G2", "애매해요", "FFF2D9", "A15C00")
xs.auto_filter.ref = f"A1:H{last}"
n_excluded = len(rows)

# ═════════════ 5. 이슈 리포트 ═════════════
iss = wb.create_sheet("이슈 리포트")
cols = ["이슈 ID", "등록일", "등록자", "화면", "주소", "환경", "유형", "심각도", "제목", "재현 단계", "기대 결과", "실제 결과",
        "스크린샷 링크", "관련 케이스·글 ID", "상태", "수정 담당", "재확인", "비고"]
header(iss, 1, cols, [9, 11, 9, 13, 24, 15, 13, 10, 36, 40, 28, 28, 20, 14, 11, 10, 10, 20])
iss.freeze_panes = "J2"
ROWS = 200
for i in range(2, ROWS + 2):
    for c in range(1, len(cols) + 1):
        cell = iss.cell(row=i, column=c)
        cell.font, cell.border = font(color=INK), BOX
        cell.alignment = CENTER if c in (1, 2, 3, 4, 6, 7, 8, 14, 15, 16, 17) else WRAP
    # 번호는 미리 박아 둔 고정 값 — 정렬하거나 줄을 지워도 다른 시트에 적은 이슈 ID 와 어긋나지 않는다
    iss.cell(row=i, column=1, value="예시" if i == 2 else f"QA-{i-2:03d}")
dv(iss, SCREENS, f"D2:D{ROWS+1}")
dv(iss, ENVS, f"F2:F{ROWS+1}")
dv(iss, ISSUE_TYPES, f"G2:G{ROWS+1}")
dv(iss, SEVERITY, f"H2:H{ROWS+1}")
dv(iss, STATUS, f"O2:O{ROWS+1}")
dv(iss, ["확인 완료", "다시 발생"], f"Q2:Q{ROWS+1}")
rng = f"H2:H{ROWS+1}"
color_rule(iss, rng, "H2", "S1 치명", "C62828", "FFFFFF")
color_rule(iss, rng, "H2", "S2 높음", "FDE7E7", "C62828")
color_rule(iss, rng, "H2", "S3 보통", "FFF2D9", "A15C00")
color_rule(iss, rng, "H2", "S4 낮음", "F2F4F6", "6B7684")
rng = f"O2:O{ROWS+1}"
color_rule(iss, rng, "O2", "재확인 완료", "E3F6EA", "1B7F3B")
color_rule(iss, rng, "O2", "수정 완료", "E8F3FF", "3182F6")
color_rule(iss, rng, "O2", "보류", "F2F4F6", "6B7684")
color_rule(iss, rng, "O2", "이슈 아님", "F2F4F6", "6B7684")
iss.auto_filter.ref = f"A1:R{ROWS+1}"
# 예시 한 줄
ex = ["", "2026-10-08", "예시", "글 상세", "/articles/41", "iPhone · Safari", "UI·레이아웃", "S3 보통",
      "(예시) 글 상세 · 기획자 관점 시트 · 닫기 버튼이 안 눌려요",
      "1. /articles/41 접속\n2. '✦ 기획자 관점 보기' 누르기\n3. 오른쪽 위 × 누르기",
      "시트가 닫혀요", "반응이 없고 바깥을 눌러야 닫혀요", "drive.google.com/...", "AR-07", "신규", "", "", "예시 줄 — 확인 후 지워 주세요"]
for c, v in enumerate(ex, start=1):
    if v:
        cell = iss.cell(row=2, column=c, value=v)
        cell.font = font(color=SUB, italic=True)

# ═════════════ 6. 현황 ═════════════
st = wb.create_sheet("현황")
st.sheet_view.showGridLines = False
st.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGH", [22, 12, 12, 12, 12, 12, 12]):
    st.column_dimensions[col].width = w
st.cell(row=2, column=2, value="QA 현황").font = font(size=18, bold=True, color=INK)
st.cell(row=3, column=2, value="다른 시트를 채우면 자동으로 바뀌어요.").font = font(color=SUB)

def table(row, title, heads, lines):
    st.cell(row=row, column=2, value=title).font = font(size=12, bold=True, color=INK)
    row += 1
    for c, h in enumerate(heads, start=2):
        cell = st.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = font(bold=True, color="FFFFFF"), HEAD, CENTER, BOX
    for line in lines:
        row += 1
        for c, v in enumerate(line, start=2):
            cell = st.cell(row=row, column=c, value=v)
            cell.font, cell.border = font(color=INK, bold=(c == 2)), BOX
            cell.alignment = Alignment(horizontal="left" if c == 2 else "center")
    return row + 2

T = "'테스트 케이스'"
n_case = len(CASES) + 1
row = table(5, "테스트 케이스", ["", "전체", "Pass", "Fail", "Block", "N/A", "진행률"], [
    [env, f"=COUNTA({T}!A2:A{n_case})",
     f'=COUNTIF({T}!{col}2:{col}{n_case},"Pass")', f'=COUNTIF({T}!{col}2:{col}{n_case},"Fail")',
     f'=COUNTIF({T}!{col}2:{col}{n_case},"Block")', f'=COUNTIF({T}!{col}2:{col}{n_case},"N/A")',
     f'=TEXT(1-COUNTIF({T}!{col}2:{col}{n_case},"미진행")/COUNTA({T}!A2:A{n_case}),"0%")']
    for env, col in (("PC", "I"), ("모바일", "J"))
])
C = "'콘텐츠 검수'"
cl = n_content + 1
row = table(row, f"콘텐츠 검수 (포함 {n_content}건)", ["", "검수함", "O 정확", "△ 어긋남", "X 틀림", "", "진행률"], [
    [name, f'=COUNTIF({C}!{col}2:{col}{cl},"?*")', f'=COUNTIF({C}!{col}2:{col}{cl},"O 정확")',
     f'=COUNTIF({C}!{col}2:{col}{cl},"△ 일부 어긋남")', f'=COUNTIF({C}!{col}2:{col}{cl},"X 틀림")', "",
     f'=TEXT(COUNTIF({C}!{col}2:{col}{cl},"?*")/{n_content},"0%")']
    for name, col in (("핵심 카피·결론", "J"), ("BEFORE→AFTER", "K"), ("더 들어가 볼까요", "L"))
] + [["경험 분류", f'=COUNTIF({C}!M2:M{cl},"?*")', f'=COUNTIF({C}!M2:M{cl},"동의")',
      f'=COUNTIF({C}!M2:M{cl},"애매함")', f'=COUNTIF({C}!M2:M{cl},"다른 경험이 맞음")+COUNTIF({C}!M2:M{cl},"해당 없음이 맞음")', "",
      f'=TEXT(COUNTIF({C}!M2:M{cl},"?*")/{n_content},"0%")']])
X = "'제외 글 검수'"
xl = n_excluded + 1
row = table(row, f"제외 글 검수 (제외 {n_excluded}건)", ["", "검수함", "맞아요", "애매해요", "포함해야 해요", "", "진행률"], [
    ["제외 판정", f'=COUNTIF({X}!G2:G{xl},"?*")', f'=COUNTIF({X}!G2:G{xl},"맞아요")', f'=COUNTIF({X}!G2:G{xl},"애매해요")',
     f'=COUNTIF({X}!G2:G{xl},"포함해야 해요")', "", f'=TEXT(COUNTIF({X}!G2:G{xl},"?*")/{n_excluded},"0%")']])
I = "'이슈 리포트'"
il = ROWS + 1
open_expr = lambda sev: (f'=COUNTIFS({I}!H3:H{il},"{sev}",{I}!O3:O{il},"<>재확인 완료",{I}!O3:O{il},"<>이슈 아님",{I}!O3:O{il},"<>보류")')
row = table(row, "이슈", ["심각도", "전체", "아직 열림", "재확인 완료", "", "", "배포 가능?"], [
    [sev, f'=COUNTIF({I}!H3:H{il},"{sev}")', open_expr(sev), f'=COUNTIFS({I}!H3:H{il},"{sev}",{I}!O3:O{il},"재확인 완료")', "", "",
     (f'=IF({open_expr(sev)[1:]}=0,"OK","막힘")' if sev in ("S1 치명", "S2 높음") else "")]
    for sev in SEVERITY
])
st.cell(row=row, column=2, value="배포 기준: S1·S2 '아직 열림'이 모두 0이면 배포해요.").font = font(color=SUB)

wb.move_sheet("현황", offset=-(len(wb.sheetnames) - 2))
wb.save(OUT)
print("saved", OUT, "cases", len(CASES), "content", n_content, "excluded", n_excluded)
